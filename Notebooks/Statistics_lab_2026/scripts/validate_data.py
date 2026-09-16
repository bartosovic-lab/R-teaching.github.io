"""Read-only reconciliation against original workbook; run from any directory."""
from pathlib import Path
import csv,json,hashlib,math,re
root=Path(__file__).resolve().parents[1]
cfg=json.loads((root/'course.json').read_text())
def rows(name): return list(csv.DictReader((root/'data'/name).open()))
core=rows('study.csv'); full=rows('investigation.csv')
assert len({r['sample_id'] for r in core})==len(core)
assert len({r['sample_id'] for r in full})==len(full)
assert sum(r['group']==cfg['group_a'] for r in core)==cfg['n_a']
assert sum(r['group']==cfg['group_b'] for r in core)==cfg['n_b']
assert (root/'data/study.csv').read_bytes()==(root/'tutorials/01_basics/data/study.csv').read_bytes()
source=next((root/'data/source').glob('*.xls*'))
expected=re.search(r'SHA-256: ([a-f0-9]{64})',(root/'data/source/PROVENANCE.md').read_text()).group(1)
assert hashlib.sha256(source.read_bytes()).hexdigest()==expected
def eq(a,b):
 if a=='NA': assert b in ('',None)
 else: assert math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-12)
if cfg['variant']=='mouse':
 import xlrd
 s=xlrd.open_workbook(source).sheet_by_index(0)
 names=s.row_values(0); lookup={s.cell_value(i,0):dict(zip(names,s.row_values(i))) for i in range(1,s.nrows)}
 assert len(full)==72
 for r in full:
  raw=lookup[r['source_record']]
  assert r['source_record']==r['sample_id']+'_1'
  for f,rawname in [('BDNF','BDNF_N'),('pCREB','pCREB_N'),('NR1','NR1_N')]: eq(r[f],raw[rawname])
  assert (r['genotype'],r['treatment'],r['learning'])==(raw['Genotype'],raw['Treatment'],raw['Behavior'])
 for r in core:
  v=next(x for x in full if x['sample_id']==r['sample_id'])
  assert v['genotype']=='Control' and v['learning']=='C/S' and r['group']==v['treatment']
  eq(r['value'],v['BDNF']); eq(r['companion'],v['pCREB'])
else:
 import openpyxl
 w=openpyxl.load_workbook(source,read_only=True,data_only=True); s=w['Figure 1G+1H']; cells=rows('cells.csv')
 assert len(cells)==46 and len(full)==26
 for r in cells:
  col=2 if r['conditioning']=='Paired' else 7
  eq(r['reversal_mV'],s.cell(int(r['source_row']),col).value)
  eq(r['resting_mV'],s.cell(int(r['source_row']),col+1).value)
 for r in full:
  group=[v for v in cells if v['sample_id']==r['sample_id']]
  assert len(group)==int(r['n_cells'])
  for f in ['reversal_mV','resting_mV']: eq(r[f],sum(float(v[f]) for v in group)/len(group))
 for r in core:
  v=next(x for x in full if x['sample_id']==r['sample_id'])
  assert v['conditioning']=='Paired' and r['group']==v['phase']
  eq(r['value'],v['reversal_mV']); eq(r['companion'],v['resting_mV'])
 protein=rows('nkcc1.csv'); s=w['Figure S1N']
 assert len(protein)==16 and len({r['sample_id'] for r in protein})==16
 for r in protein:
  n=int(r['source_row']); col=4 if r['day']=='1' else 7
  eq(r['NKCC1_GAPDH'],s.cell(n,col).value)
  assert r['sample_id']==s.cell(n,col-1).value
  # Independently locate the ID in raw band-intensity rows and verify the supplied ratio.
  rawrow=next(i for i in list(range(4,12))+list(range(13,21)) if s.cell(i,1).value==r['sample_id'])
  eq(r['NKCC1_GAPDH'],s.cell(rawrow,10).value)
  assert math.isclose(float(r['NKCC1_GAPDH']),s.cell(rawrow,5).value/s.cell(rawrow,3).value,rel_tol=1e-8)
  assert r['blot_batch']==('February' if rawrow<12 else 'March')
print('PASS',cfg['variant'],': source checksum; exact source values; selection/aggregation; sample counts; tutorial data copy')
